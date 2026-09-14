# -*- coding: utf-8 -*-
"""Construye la version APA 7 v2 de la tesis Agroplaga AI a partir de la version 'APA7 editado'.
Uso: python construir_docx.py <fig_dir> <salida.docx>"""
import sys, os, re, copy, itertools
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph
from docx.table import Table
from PIL import Image

FIG = sys.argv[1]
DST = sys.argv[2]
SRC = r"C:\Users\cesar\Desktop\Proyecto de graduacion\ProyectoGaduacionAP\Proyecto de Graducion Daniel Cifuentes cap6 - APA7 editado.docx"
REPO = "https://github.com/Cifzu233/ProyectoGaduacionAP"

doc = Document(SRC)
P = list(doc.paragraphs)          # P[i-1] = parrafo i (numeracion del analisis)
T = list(doc.tables)
def p(i): return P[i - 1]
CENTER, LEFT, RIGHT, JUST = WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.RIGHT, WD_ALIGN_PARAGRAPH.JUSTIFY

# ------------------------------------------------------------------ helpers
CTRL_RE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
def clear_runs(par):
    """Elimina los runs. Devuelve True si el parrafo contenia un salto de pagina (para conservarlo)."""
    had_pb = any(br.get(qn("w:type")) == "page" for br in par._p.findall(".//" + qn("w:br")))
    for child in list(par._p):
        if child.tag != qn("w:pPr"):
            par._p.remove(child)
    return had_pb

def add_runs(par, text, bold=False, italic=False, size=None):
    text = CTRL_RE.sub("", text)
    parts = text.split("\n")
    for k, part in enumerate(parts):
        r = par.add_run(part)
        r.bold = bold or None; r.italic = italic or None
        r.font.name = "Times New Roman"; r.font.size = Pt(size or 12)
        if k < len(parts) - 1: r.add_break()
    return par

def set_text(par, text, bold=False, italic=False, align=None, style=None):
    had_pb = clear_runs(par); add_runs(par, text, bold, italic)
    if had_pb: page_break_after(par)
    if style: par.style = doc.styles[style]
    if align is not None: par.alignment = align
    return par

def par_after(anchor, text="", style=None, align=None, bold=False, italic=False, size=None):
    new_p = OxmlElement("w:p"); anchor._p.addnext(new_p)
    np_ = Paragraph(new_p, anchor._parent)
    if style: np_.style = doc.styles[style]
    if text: add_runs(np_, text, bold, italic, size)
    if align is not None: np_.alignment = align
    return np_

def chain(anchor, items):
    """items: lista de (texto, style) o (texto, style, dict). Devuelve el ultimo parrafo."""
    cur = anchor
    for it in items:
        text, style = it[0], it[1]; extra = it[2] if len(it) > 2 else {}
        cur = par_after(cur, text, style=style, **extra)
    return cur

def delete_par(par):
    el = par._p; el.getparent().remove(el)

def add_field(par, instr, placeholder=""):
    def fld(t):
        r = par.add_run(); f = OxmlElement("w:fldChar"); f.set(qn("w:fldCharType"), t); r._r.append(f)
    fld("begin")
    r = par.add_run(); it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = f" {instr} "; r._r.append(it)
    fld("separate")
    rr = par.add_run(placeholder); rr.font.name = "Times New Roman"; rr.font.size = Pt(12)
    fld("end")

def page_break_after(par):
    r = par.add_run(); r.add_break(__import__("docx").enum.text.WD_BREAK.PAGE)

def fig_size(path, max_w=6.0, max_h=6.8):
    with Image.open(path) as im:
        w, h = im.size
    ratio = h / w
    wi = max_w
    if wi * ratio > max_h: wi = max_h / ratio
    return Inches(wi)

def picture_after(anchor, path, max_w=6.0, max_h=6.8):
    np_ = par_after(anchor, "", align=CENTER)
    np_.add_run().add_picture(path, width=fig_size(path, max_w, max_h))
    return np_

SEQ_MARK = "SEQ|"
def caption_after(anchor, kind, text):
    """Leyenda con campo SEQ (numeracion automatica): 'Figura N. texto'. Se marca y se construye en la pasada final."""
    np_ = par_after(anchor, f"{SEQ_MARK}{kind}|{text}", align=CENTER)
    return np_

def replace_picture(par, path, max_w=6.0, max_h=6.8):
    clear_runs(par); par.alignment = CENTER
    par.add_run().add_picture(path, width=fig_size(path, max_w, max_h))

def set_cell(cell, text, bold=False, size=11):
    par = cell.paragraphs[0]
    for extra in cell.paragraphs[1:]: delete_par(extra)
    clear_runs(par); add_runs(par, text, bold=bold, size=size)
    par.paragraph_format.line_spacing = 1.0

def fill_table(tbl, rows, header_bold=True, size=11):
    while len(tbl.rows) < len(rows): tbl.add_row()
    while len(tbl.rows) > len(rows):
        tr = tbl.rows[-1]._tr; tr.getparent().remove(tr)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            set_cell(tbl.rows[ri].cells[ci], val, bold=(header_bold and ri == 0), size=size)

def table_after(anchor, template_tbl, rows, size=11):
    new_tbl = copy.deepcopy(template_tbl._tbl); anchor._p.addnext(new_tbl)
    t = Table(new_tbl, anchor._parent); fill_table(t, rows, size=size)
    return t

def para_after_table(tbl):
    return par_after_el(tbl._tbl)

def par_after_el(el, text="", style=None, align=None, italic=False):
    new_p = OxmlElement("w:p"); el.addnext(new_p)
    np_ = Paragraph(new_p, doc._body)
    if style: np_.style = doc.styles[style]
    if text: add_runs(np_, text, italic=italic)
    if align is not None: np_.alignment = align
    return np_

H1, H2, H3, H4, N = "Heading 1", "Heading 2", "Heading 3", "Heading 4", "Normal"

# ------------------------------------------------------------------ 1. mojibake
VOWELS = "áéíóúñüÁÉÍÓÚÑ"
vocab = set()
for par in P:
    if "?" in par.text: continue
    for w in re.findall(r"[A-Za-zÁÉÍÓÚÑÜáéíóúñü]+", par.text): vocab.add(w.lower())
MANUAL = {"?ndice": "Índice", "?nicamente": "únicamente", "T?cnica": "Técnica", "L?mites": "Límites", "Cap?tulo": "Capítulo",
          "env?a": "envía", "despu?s": "después", "m?s": "más", "tem?tica": "temática", "tem?tico": "temático", "l?gica": "lógica",
          "L?gica": "Lógica", "agronom?a": "agronomía", "agron?mica": "agronómica", "agron?mico": "agronómico", "m?dulos": "módulos",
          "m?dulo": "módulo", "valid?": "validó", "construy?": "construyó", "desarroll?": "desarrolló", "permiti?": "permitió",
          "im?genes": "imágenes", "gr?ficos": "gráficos", "melocot?n": "melocotón", "diagn?stico": "diagnóstico", "informaci?n": "información"}
def fix_token(tok):
    if tok in MANUAL: return MANUAL[tok]
    core = re.match(r"^(\W*)(.*?)(\W*)$", tok, re.S)
    pre, mid, suf = core.groups()
    if mid in MANUAL: return pre + MANUAL[mid] + suf
    idx = [i for i, ch in enumerate(mid) if ch == "?"]
    if not idx or len(idx) > 2: return tok
    for combo in itertools.product(VOWELS, repeat=len(idx)):
        cand = list(mid)
        for i, ch in zip(idx, combo): cand[i] = ch
        cand = "".join(cand)
        if cand.lower() in vocab: return pre + cand + suf
    # heuristica: terminaciones frecuentes
    for pat, rep in ((r"ci\?n$", "ción"), (r"si\?n$", "sión"), (r"\?n$", "ón"), (r"^\?", "ú")):
        if re.search(pat, mid): return pre + re.sub(pat, rep, mid) + suf
    return tok

def fix_mojibake(par):
    txt = par.text
    if "?" not in txt or "http" in txt: return
    if not re.search(r"\w\?\w|\?\w", txt): return
    lines = txt.split("\n"); out = []
    for ln in lines:
        if re.match(r"^\?\s", ln): ln = "• " + ln[2:]
        out.append(" ".join(fix_token(t) if "?" in t else t for t in ln.split(" ")))
    fixed = "\n".join(out)
    if fixed != txt:
        bold = any(r.bold for r in par.runs); italic = any(r.italic for r in par.runs)
        set_text(par, fixed, bold=bold and par.style.name == N, italic=italic and par.style.name == N)

for par in P: fix_mojibake(par)

# ------------------------------------------------------------------ 2. portada / indices / intro
set_text(p(2), "FACULTAD DE INGENIERÍA EN SISTEMAS DE INFORMACIÓN Y CIENCIAS DE LA", align=CENTER)
set_text(p(3), "COMPUTACIÓN", align=CENTER)
set_text(p(11), "DANIEL ANÍBAL CIFUENTES RAMÍREZ", align=CENTER)

set_text(p(31), "Índice General", bold=True, align=CENTER)
toc = par_after(p(31), ""); add_field(toc, 'TOC \\o "1-3" \\h \\z \\u', "[Índice: haga clic derecho → Actualizar campos]"); page_break_after(toc)
set_text(p(32), "Índice de Figuras", bold=True, align=CENTER)
tf = par_after(p(32), ""); add_field(tf, 'TOC \\h \\z \\c "Figura"', "[Índice de figuras: actualizar campos]"); page_break_after(tf)
set_text(p(33), "Índice de Tablas", bold=True, align=CENTER)
tt = par_after(p(33), ""); add_field(tt, 'TOC \\h \\z \\c "Tabla"', "[Índice de tablas: actualizar campos]"); page_break_after(tt)
p(37).style = doc.styles[H1]; p(37).alignment = CENTER

# ------------------------------------------------------------------ 3. capitulo 1
set_text(p(82), "Implementar un nodo IoT de bajo costo para monitorear las condiciones ambientales de la parcela y captar imágenes del cultivo. "
         "El nodo está formado por un módulo ESP32-CAM (microcontrolador con Wi-Fi y cámara OV2640) y un sensor DHT22 de temperatura y humedad "
         "del aire, alimentados por un punto de carga solar (panel fotovoltaico, módulo de carga y batería) que permite su operación autónoma "
         "en campo. El nodo transmite las lecturas y el video por Wi-Fi al servidor de la aplicación.")
set_text(p(83), "Integrar un servicio de inteligencia artificial para la identificación automática de plagas a partir de las imágenes captadas "
         "por la cámara y del contexto ambiental de la parcela. El sistema utiliza un modelo multimodal de visión accesible mediante la API "
         "de OpenAI (OpenAI, 2024), al que se le envía la imagen junto con las últimas lecturas y umbrales de la parcela, y que devuelve un "
         "diagnóstico estructurado con la plaga detectada, el nivel de confianza, la severidad y recomendaciones de manejo integrado. "
         "Este enfoque evita entrenar un modelo desde cero y permite un análisis en tiempo real con recursos limitados.")
set_text(p(84), "Diseñar una aplicación web que proporcione a los agricultores un entendimiento detallado sobre el estado de sus cultivos y "
         "recomendaciones de manejo contra las plagas. La solución es accesible desde cualquier navegador y se desarrolla con React y Vite para "
         "la interfaz, Node.js con Express para el servidor y MySQL como base de datos. La interfaz se diseña en Figma priorizando la "
         "usabilidad y una experiencia visual clara para los agricultores.")
set_text(p(94), "¿Cómo puede diseñarse la aplicación web para que sea accesible y útil para los agricultores?")
set_text(p(101), "La tecnología necesaria para este proyecto es accesible y factible de implementar en el contexto de FRUTAGRU. El módulo "
         "ESP32-CAM integra en una sola placa el microcontrolador, la conectividad Wi-Fi y una cámara, y el sensor DHT22 mide temperatura y "
         "humedad con precisión suficiente para el monitoreo agrícola; ambos tienen un costo reducido y amplia documentación. Un punto de "
         "carga solar permite instalar el nodo en la parcela sin depender de la red eléctrica. Las imágenes y lecturas se envían a un "
         "servidor web que las integra con un servicio de inteligencia artificial para analizar patrones en tiempo real, y los agricultores "
         "reciben alertas y recomendaciones desde la aplicación web de manera sencilla y rápida.")
set_text(p(103), "Como responsable de este proyecto, he considerado únicamente los costos relacionados con la adquisición del hardware del "
         "nodo de campo, ya que representa el componente esencial del monitoreo. Un módulo ESP32-CAM tiene un costo aproximado de Q120, "
         "un sensor DHT22 de Q70 y los componentes del punto de carga solar (panel de 6 V, módulo de carga TP4056, batería 18650 y "
         "elevador de voltaje) suman cerca de Q200. En total, un nodo completo se estima en Q400 – Q450, además del hosting y dominio de la "
         "aplicación web. A largo plazo, este costo se compensará con la reducción en pérdidas por plagas y el menor uso de pesticidas.")
set_text(p(104), "Operativa: Uno de los retos principales es asegurar que el sistema sea fácil de utilizar por los agricultores. Por esta "
         "razón, la aplicación web se diseñó con una interfaz intuitiva, video en vivo de la parcela y alertas visuales, de manera que los "
         "productores puedan recibir información clara sobre la detección de plagas y las acciones recomendadas. Además, se impartirán "
         "capacitaciones dentro de FRUTAGRU para garantizar que los agricultores se familiaricen con el uso del sistema y puedan integrarlo "
         "en su rutina de trabajo sin dificultades.")
# 1.5 general / especificas
g = par_after(p(86), "1.5.1 Pregunta general", style=H3)
g = par_after(g, "¿Cómo puede desarrollarse un sistema inteligente basado en sensores IoT, video en tiempo real e inteligencia artificial "
              "que permita detectar de forma temprana las plagas en los cultivos de melocotón de FRUTAGRU y apoyar la toma de decisiones "
              "de los agricultores?")
par_after(g, "1.5.2 Preguntas específicas", style=H3)

# ------------------------------------------------------------------ 4. capitulo 2
m = par_after(p(217), "2.8.2. Materiales", style=H3)
m = chain(m, [
    ("• Computadora portátil: utilizada para el diseño, desarrollo, pruebas y documentación del sistema.", N),
    ("• Nodo de campo: módulo ESP32-CAM (AI-Thinker, sensor OV2640), sensor DHT22, panel solar de 6 V, módulo de carga TP4056, batería Li-ion 18650, elevador de voltaje a 5 V, caja estanca y cableado.", N),
    ("• Adaptador USB-serie para programar el ESP32-CAM y multímetro para verificar la alimentación.", N),
    ("• Conexión a internet y red Wi-Fi de 2.4 GHz en la parcela de prueba.", N),
    ("• Software y herramientas: Visual Studio Code, Node.js, React, MySQL (en Docker), Arduino IDE, Lucidchart, Figma, Postman y el servicio de IA de OpenAI.", N),
])
set_text(p(218), "2.8.3. Financieros", style=H3)
set_text(p(219), "2.8.3.1 Implementación del proyecto", style=H4)
set_text(p(222), "2.8.3.2 Costo del desarrollo del proyecto", style=H4)
fill_table(T[1], [
    ["Cantidad", "Descripción", "Costo unitario (Q.)", "Costo total (Q.)"],
    ["1", "Módulo ESP32-CAM (AI-Thinker, OV2640) con adaptador USB", "Q120.00", "Q120.00"],
    ["1", "Sensor de temperatura y humedad DHT22", "Q70.00", "Q70.00"],
    ["1", "Panel solar 6 V / 5 W", "Q95.00", "Q95.00"],
    ["1", "Módulo de carga TP4056 con protección", "Q15.00", "Q15.00"],
    ["1", "Batería Li-ion 18650 (3000 mAh) con portapilas", "Q45.00", "Q45.00"],
    ["1", "Elevador de voltaje 5 V (MT3608)", "Q20.00", "Q20.00"],
    ["1", "Caja estanca, cables y resistencia pull-up", "Q60.00", "Q60.00"],
    ["1", "Hosting web (mensual)", "Q43.00", "Q43.00"],
    ["1", "Dominio web (anual)", "Q200.00", "Q200.00"],
    ["1", "Impresiones y papelería", "Q100.00", "Q100.00"],
    ["", "", "Total", "Q768.00"],
])
fill_table(T[2], [
    ["Cantidad", "Descripción", "Costo por hora (Q.)", "Horas requeridas", "Total (Q.)"],
    ["1", "Desarrollador fullstack (Daniel Cifuentes)", "Q. 0.00", "480", "Q. 0.00"],
    ["1", "Integrador de hardware (ESP32-CAM, DHT22 y energía solar)", "Q. 0.00", "60", "Q. 0.00"],
    ["1", "Integración del servicio de IA de visión", "Q. 0.00", "120", "Q. 0.00"],
    ["", "", "", "Total", "Q. 0.00"],
])

# ------------------------------------------------------------------ 5. capitulo 3
set_text(p(299), "DHT22: sensor digital de temperatura y humedad del aire con precisión de ±0,5 °C y ±2 % HR (Kumar & Lee, 2019). Se conecta al "
         "ESP32-CAM mediante un solo pin de datos y es adecuado para caracterizar las condiciones que favorecen el desarrollo de plagas como "
         "pulgones y ácaros.")
delete_par(p(300))
set_text(p(301), "ESP32-CAM: módulo basado en el microcontrolador ESP32 con Wi-Fi y Bluetooth incorporados, memoria PSRAM y cámara OV2640 de "
         "2 megapíxeles (Espressif Systems, 2023). Puede servir video MJPEG, capturar fotografías JPEG y leer sensores digitales, lo que lo "
         "convierte en un nodo de campo completo de bajo costo.")
set_text(p(302), "Punto de carga solar: conjunto formado por un panel fotovoltaico de 6 V, un módulo de carga TP4056 con protección de batería, "
         "una batería Li-ion 18650 y un elevador de voltaje a 5 V. Permite alimentar el nodo de forma autónoma en parcelas sin acceso a la red "
         "eléctrica.")
s = par_after(p(305), "3.4.4. Alimentación autónoma con energía solar", style=H3)
s = chain(s, [
    ("Los nodos IoT instalados en campo rara vez disponen de una toma eléctrica cercana, por lo que la alimentación autónoma es un requisito "
     "práctico. Un sistema fotovoltaico pequeño se compone de un panel que convierte la radiación solar en corriente continua, un circuito de "
     "carga que regula la corriente hacia una batería de litio y protege contra sobrecarga y descarga profunda, y un regulador que entrega el "
     "voltaje estable que requiere el microcontrolador.", N),
    ("En Agroplaga AI el ESP32-CAM consume entre 180 y 300 mA cuando transmite video y menos de 80 mA en reposo, mientras que el DHT22 requiere "
     "menos de 2 mA. Con un panel de 5 W y una batería de 3000 mAh el nodo funciona de forma continua durante el día y conserva energía para "
     "seguir enviando lecturas durante la noche, activando el video únicamente cuando un usuario lo solicita desde la aplicación web.", N),
])
set_text(p(311), "Las Redes Neuronales Convolucionales (CNN) son la base para análisis de imágenes de hojas y frutos. Arquitecturas como "
         "ResNet-50 o MobileNet, entrenadas con transfer learning, han logrado precisiones superiores al 92 % en detección de plagas en manzano, "
         "resultados transferibles al melocotón (Zhang et al., 2019; Kamilaris & Prenafeta-Boldú, 2018). Más recientemente, los modelos "
         "multimodales de gran escala, como GPT-4o, combinan visión y lenguaje y permiten describir una imagen, identificar organismos y "
         "generar recomendaciones en lenguaje natural sin necesidad de un conjunto de entrenamiento propio (OpenAI, 2024). Agroplaga AI adopta "
         "este segundo enfoque, accediendo al modelo a través de una API y enviándole, junto con la imagen, el contexto ambiental de la parcela.")
set_text(p(313), "La clasificación automática reduce errores humanos y acelera la respuesta. Los sistemas embebidos con TensorFlow Lite "
         "permiten inferencia offline con latencias menores a 200 ms, mientras que los servicios en la nube tardan entre 2 y 8 segundos por "
         "imagen pero ofrecen mayor capacidad de generalización y explicaciones en lenguaje natural (Li et al., 2020). Para un cultivo en el que "
         "las plagas evolucionan en días, esta latencia es aceptable y se compensa con la calidad del diagnóstico y las recomendaciones.")
set_text(p(316), "Una aplicación web se ejecuta en el navegador sin instalación, se actualiza de forma centralizada y funciona en computadoras "
         "y teléfonos, lo que facilita su adopción y mantenimiento en comunidades rurales (Lehrig, 2021). Además, un servidor web puede "
         "concentrar la conexión con los dispositivos de campo y retransmitir su video a varios usuarios a la vez, algo que un dispositivo "
         "IoT por sí solo no podría soportar.")
set_text(p(324), "El frontend se construye con React y Vite como aplicación de una sola página, y se comunica mediante HTTP con un backend "
         "en Node.js y Express que expone una API REST. La información se almacena en MySQL, accedida mediante Sequelize. El servidor abre "
         "una única conexión con el ESP32-CAM, trocea el flujo MJPEG y lo retransmite a los navegadores conectados, y envía las imágenes al "
         "servicio de visión de OpenAI para su diagnóstico. El firmware del nodo se desarrolla en Arduino para el ESP32.")
a = par_after(p(324), "3.6.4. Arquitectura del sistema Agroplaga AI", style=H3)
chain(a, [
    ("La arquitectura del sistema se organiza en tres capas. La capa de campo está formada por el nodo ESP32-CAM con el sensor DHT22 y su "
     "punto de carga solar, que publica lecturas ambientales y video por Wi-Fi. La capa de servidor, implementada en Node.js, recibe las "
     "lecturas, evalúa umbrales y genera alertas, mantiene en memoria el último cuadro de cada cámara para retransmitirlo, y coordina el "
     "análisis de imágenes con el servicio de inteligencia artificial, guardando cada detección en MySQL. La capa de presentación es la "
     "aplicación web en React, desde la que el agricultor consulta el panel ambiental, el video en vivo, el historial de detecciones y el "
     "asistente agronómico. El diagrama de arquitectura se presenta en el capítulo 5 (sección 5.2.2).", N),
])

# ------------------------------------------------------------------ 6. capitulo 4
a = par_after(p(336), "4.2 Análisis de requerimientos", style=H2)
a = chain(a, [
    ("4.2.1 Propósito del sistema", H3),
    ("El propósito de Agroplaga AI es brindar a los productores de melocotón de FRUTAGRU una herramienta accesible que integre el monitoreo "
     "ambiental de la parcela, el video en vivo del cultivo y el diagnóstico de plagas mediante inteligencia artificial, de modo que puedan "
     "detectar problemas fitosanitarios a tiempo y decidir tratamientos con información objetiva.", N),
    ("4.2.2 Alcance del sistema", H3),
    ("El sistema está disponible como aplicación web para computadoras y teléfonos. Recibe temperatura y humedad del sensor DHT22 y el video "
     "del ESP32-CAM instalados en la parcela, genera alertas por umbrales, permite capturar y analizar imágenes con IA de forma manual o "
     "automática, administra parcelas, fichas de plagas y actividades agrícolas, y ofrece un asistente conversacional restringido a la "
     "agronomía del melocotón. El alcance se limita a las parcelas piloto de FRUTAGRU con un nodo de campo alimentado por energía solar.", N),
    ("4.2.3 Restricciones del sistema", H3),
    ("• El nodo requiere cobertura Wi-Fi de 2.4 GHz en la parcela y el servidor debe poder alcanzar la dirección IP del ESP32-CAM o, en su "
     "defecto, el nodo debe enviar las fotografías al servidor (modo push).\n"
     "• El ESP32-CAM admite uno o dos clientes de video simultáneos; por ello el servidor centraliza la conexión y la retransmite.\n"
     "• El diagnóstico por IA depende de un servicio externo con costo por uso, por lo que el análisis automático tiene un intervalo mínimo de 60 segundos.\n"
     "• La precisión del diagnóstico depende de la iluminación, el enfoque y la distancia de la cámara al follaje.", N),
    ("4.2.4 Características de los usuarios", H3),
    ("Los usuarios principales son los agricultores de FRUTAGRU, con distintos niveles de familiaridad tecnológica, y los técnicos de la "
     "asociación que apoyan el manejo de plagas. También se contempla un rol administrador encargado de registrar parcelas, cámaras y fichas "
     "de plagas. Por ello la interfaz utiliza lenguaje sencillo, tarjetas visuales, colores de severidad y la menor cantidad de texto posible.", N),
    ("4.2.5 Interfaces del sistema", H3),
    ("• Interfaz de usuario: aplicación web con panel ambiental, página de cámara en vivo, gestión de plagas, parcelas, seguimiento y asistente agronómico.\n"
     "• Interfaz de hardware: el ESP32-CAM lee el DHT22 por el pin GPIO 13 y se alimenta a 5 V desde el punto de carga solar.\n"
     "• Interfaz de comunicación: el nodo envía lecturas JSON a POST /api/readings y sirve video MJPEG en el puerto 81; el servidor expone una "
     "API REST en /api y consume la API de OpenAI por HTTPS.", N),
])
set_text(p(337), "4.3 Requerimientos específicos", style=H2)
set_text(p(338), "4.3.1 Requerimientos funcionales", style=H3)
set_text(p(341), "RF3. Visualización de condiciones ambientales: El sistema debe mostrar la temperatura y la humedad del aire medidas por el DHT22 "
         "de cada parcela, con la hora de la última lectura.")
set_text(p(347), "RF9. Reportes: El sistema debe permitir exportar el historial de actividades y detecciones en formato CSV para su análisis.")
r = chain(p(347), [
    ("RF10. Video en vivo: El sistema debe mostrar en la aplicación web el video del ESP32-CAM instalado en la parcela, indicando si la cámara está en línea.", N),
    ("RF11. Análisis automático: El sistema debe permitir programar el análisis periódico de la imagen de la cámara por IA y registrar cada detección con su severidad.", N),
    ("RF12. Asistente agronómico: El sistema debe responder consultas sobre el manejo del melocotón usando el contexto de la parcela y rechazar temas ajenos al cultivo.", N),
])
set_text(p(348), "4.3.2 Requerimientos no funcionales", style=H3)
set_text(p(350), "RNF2. Rendimiento: El tiempo de respuesta del sistema al analizar una imagen no debe exceder los 10 segundos y el video debe mostrarse con una latencia menor a 2 segundos.")
q = par_after(p(354), "4.3.3 Requerimientos del sistema", style=H3)
q = chain(q, [
    ("• Nodo de campo: ESP32-CAM AI-Thinker con PSRAM, sensor DHT22, panel solar de 6 V, módulo TP4056, batería 18650 y elevador de 5 V.\n"
     "• Servidor: Node.js 18 o superior, MySQL 8, acceso a internet para el servicio de IA y una clave de API de OpenAI.\n"
     "• Cliente: navegador moderno (Chrome, Edge o Firefox) en computadora o teléfono.", N),
    ("4.3.4 Requerimientos del usuario", H3),
    ("El usuario debe poder registrar su parcela y su cámara con la dirección del ESP32-CAM, ver el estado de la cámara y del sensor, "
     "solicitar un análisis con un solo clic, interpretar el resultado mediante colores de severidad y recomendaciones, y consultar el "
     "historial sin necesidad de conocimientos técnicos.", N),
    ("4.4 Requerimientos priorizados", H2),
    ("1. Recepción de lecturas del DHT22 y alertas por umbral.\n2. Video en vivo del ESP32-CAM en la aplicación web.\n"
     "3. Captura y análisis de imágenes con IA (manual y automático).\n4. Registro de detecciones, parcelas, plagas y actividades.\n"
     "5. Asistente agronómico con contexto de parcela.\n6. Exportación de historial en CSV.", N),
])
set_text(p(356), "4.5 Herramientas y tecnologías de desarrollo", style=H2)
set_text(p(358), "Lucidchart: Diagramas UML (casos de uso, clases, secuencia, arquitectura).")
set_text(p(361), "MySQL: Base de datos relacional para parcelas, lecturas, alertas, plagas, cámaras y detecciones.")
set_text(p(362), "Arduino IDE: Desarrollo del firmware del ESP32-CAM (video MJPEG, captura y lectura del DHT22).")
set_text(p(363), "React y Vite: Desarrollo del frontend de la plataforma; Node.js con Express para la API REST.")
set_text(p(364), "Docker: Contenedor de MySQL para el entorno de desarrollo y pruebas.")
chain(p(364), [("API de OpenAI: Servicio de inteligencia artificial de visión para el diagnóstico de plagas a partir de imágenes.", N),
               ("Postman y jest: Pruebas de la API y pruebas automatizadas del backend.", N)])
set_text(p(365), "4.6 Validación de requerimientos y recomendaciones de expertos", style=H2)
set_text(p(366), "Los requerimientos se validaron con el representante de FRUTAGRU y con los ingenieros invitados a la presentación del "
         "anteproyecto. La principal recomendación fue que la detección de plagas se apoyara en una cámara en tiempo real y no solo en los "
         "sensores de humedad y temperatura, y que el prototipo fuera autónomo energéticamente para poder instalarse en la parcela. Por "
         "razones de presupuesto se acordó implementar un único nodo de prueba con ESP32-CAM, DHT22 y punto de carga solar, dejando la "
         "ampliación a más parcelas como trabajo futuro sujeto al apoyo económico de la asociación.")
set_text(p(367), "4.7 Acceso al documento ERS completo", style=H2)

# ------------------------------------------------------------------ 7. capitulo 5
set_text(p(388), "Capítulo 5 – Diagramas UML", style=H1, align=CENTER)
par_after(p(388), "El presente capítulo modela el comportamiento y la estructura de Agroplaga AI mediante diagramas UML. Los diagramas de "
          "comportamiento describen cómo interactúan el agricultor, el técnico, el administrador y el nodo ESP32-CAM con el sistema, y los "
          "diagramas estructurales muestran las clases, la arquitectura y los componentes que lo integran. Al final se presentan las "
          "interfaces de la aplicación desarrollada.")
set_text(p(389), "5.1 Diagramas de comportamiento", style=H2)
replace_picture(p(391), os.path.join(FIG, "casos_uso.png"))
set_text(p(393), "Figura 6. Diagrama de casos de uso: interacciones de los actores (agricultor, técnico, administrador y nodo ESP32-CAM) con Agroplaga AI. (Fuente: elaboración propia)")
for ti in range(3, 8):
    for row in T[ti].rows:
        for cell in row.cells:
            for par in cell.paragraphs:
                t = par.text
                t2 = t.replace("PDF", "CSV").replace("servidor IA", "servicio de IA").replace("humedad del suelo", "humedad del aire")
                if t2 != t: set_text(par, t2)
tbl9 = table_after(p(412), T[3], [
    ["Campo", "Descripción"], ["Nombre de caso de uso", "Ver video en vivo y analizar con IA"], ["Área", "Módulo de cámara"],
    ["Identificador", "6"], ["Actores", "Agricultor, Técnico"],
    ["Descripción", "El usuario selecciona una cámara registrada, observa el video en vivo retransmitido por el servidor y solicita un análisis del cuadro actual; el sistema guarda la detección con su severidad y recomendaciones."],
    ["Evento desencadenador", "El usuario abre la página Cámara y pulsa «Analizar con IA», o el análisis automático alcanza su intervalo."],
    ["Precondiciones", "Cámara registrada y en línea; servidor con acceso al servicio de IA."],
    ["Postcondiciones", "Detección registrada en el historial; alerta visible en el panel si requiere acción."],
    ["Requerimiento cumplido", "RF10, RF11"], ["Prioridad", "Alta"], ["Riesgo", "Medio"],
])
caption_after(para_after_table(tbl9), "Tabla", "Caso de uso 6: Ver video en vivo y analizar con IA. (Fuente: elaboración propia)")
set_text(p(427), "Representan el orden de interacción entre los actores, la aplicación web, el servidor, el servicio de inteligencia artificial y la base de datos.")
set_text(p(428), "Secuencia 1: Captura y análisis de imagen con alerta.\nSecuencia 2: Registro de lecturas del DHT22 y alerta por umbral.\nSecuencia 3: Video en vivo y análisis automático.")
delete_par(p(429))
replace_picture(p(431), os.path.join(FIG, "sec_imagen.png"))
set_text(p(432), "Figura 9. Diagrama de secuencia: captura y análisis de imagen con alerta. (Fuente: elaboración propia)")
replace_picture(p(433), os.path.join(FIG, "sec_sensores.png"))
set_text(p(434), "Figura 10. Diagrama de secuencia: registro de lecturas del DHT22 y alerta por umbral. (Fuente: elaboración propia)")
replace_picture(p(437), os.path.join(FIG, "sec_video.png"))
set_text(p(438), "Figura 11. Diagrama de secuencia: video en vivo mediante retransmisión MJPEG y análisis automático. (Fuente: elaboración propia)")
set_text(p(441), "Muestra los estados por los que pasa una imagen desde que la captura el ESP32-CAM hasta que su diagnóstico queda registrado y, si corresponde, genera una alerta.")
replace_picture(p(442), os.path.join(FIG, "estados.png"), max_h=7.5)
set_text(p(444), "Figura 12. Diagrama de estados: ciclo de vida de una imagen en Agroplaga AI. (Fuente: elaboración propia)")
replace_picture(p(447), os.path.join(FIG, "clases.png"), max_h=7.2)
set_text(p(448), "Figura 13. Diagrama de clases: entidades del sistema (Parcela, Plot, TipoSensor, LecturaSensor, Umbral, Alerta, Plaga, Actividad, Cámara y Detección) y su relación con el nodo ESP32-CAM y el servicio de visión. (Fuente: elaboración propia)")
arq = par_after(p(448), "5.2.2 Diagrama de arquitectura", style=H3)
arq = par_after(arq, "Presenta las tres capas del sistema: el nodo de campo (ESP32-CAM, DHT22 y punto de carga solar), el servidor Node.js con "
                "sus módulos de lecturas, retransmisión de video, análisis de IA y base de datos MySQL, y los servicios externos y el navegador del usuario.")
arq = picture_after(arq, os.path.join(FIG, "arquitectura.png"))
caption_after(arq, "Figura", "Diagrama de arquitectura del sistema Agroplaga AI. (Fuente: elaboración propia)")
set_text(p(451), "5.2.3 Diagrama de componentes", style=H3)
set_text(p(453), "Frontend React Web: páginas de Dashboard, Plagas, Seguimiento, Parcelas, Cámara en vivo y asistente.")
set_text(p(454), "Backend API (Node.js y Express): rutas de negocio y módulo de cámaras (FrameHub, parser MJPEG, conector upstream y scheduler).")
set_text(p(455), "Módulo de IA: construcción del prompt, llamada al servicio de visión de OpenAI y normalización del diagnóstico.")
set_text(p(456), "Base de datos MySQL: parcelas, lecturas, alertas, plagas, actividades, cámaras y detecciones.")
chain(p(456), [("Nodo ESP32-CAM + DHT22: firmware Arduino que sirve el video, captura fotografías y envía las lecturas.", N),
               ("Simulador de cámara: script Node.js que reproduce un ESP32-CAM para pruebas sin hardware.", N)])
replace_picture(p(458), os.path.join(FIG, "componentes.png"))
set_text(p(459), "Figura 14. Diagrama de componentes de Agroplaga AI. (Fuente: elaboración propia)")
ui = par_after(p(459), "5.3 Interfaces de la aplicación", style=H2)
ui = par_after(ui, "A continuación se presentan las interfaces implementadas de Agroplaga AI, que materializan los prototipos diseñados en Figma durante la fase de elaboración.")
UI = [
    ("5.3.1 Panel de monitoreo ambiental", "ui_dashboard.png", "Panel de monitoreo con temperatura y humedad del DHT22, última detección de cámara y alertas inteligentes.",
     "• Selección de la parcela de sensores y actualización automática cada 20 segundos.\n• Tarjetas de temperatura y humedad con la hora de la última lectura.\n• Tarjeta con la última detección de la cámara y enlace al video en vivo.\n• Alertas inteligentes por condiciones ambientales y por detecciones que requieren acción."),
    ("5.3.2 Cámara en vivo", "ui_camara.png", "Página de cámara en vivo con video MJPEG, estado de la cámara, análisis con IA y análisis automático.",
     "• Selección de la cámara registrada y estado en línea con cuadros por segundo.\n• Botones para pausar, reconectar, capturar y analizar con IA.\n• Selector de intervalo de análisis automático.\n• Historial de detecciones recientes con severidad y confianza."),
    ("5.3.3 Gestión de plagas", "ui_plagas.png", "Listado y edición de fichas de plagas del melocotón con imagen de referencia.", "• Consulta de nombre, nombre científico, síntomas y tratamiento.\n• Edición en línea y eliminación de fichas.\n• Registro de nuevas plagas con carga de imagen."),
    ("5.3.4 Seguimiento de actividades", "ui_seguimiento.png", "Registro de actividades agrícolas por parcela con historial, gráfica y exportación a CSV.", "• Registro de fumigaciones, podas y observaciones con fotografía.\n• Filtro por parcela y gráfica mensual de actividades.\n• Exportación del historial en formato CSV."),
    ("5.3.5 Administración de parcelas", "ui_parcelas.png", "Administración de parcelas con ubicación, cultivo, superficie y fecha de siembra.", "• Alta, edición y baja de parcelas.\n• Carga de imagen de la parcela.\n• Asociación de las parcelas con cámaras y actividades."),
    ("5.3.6 Asistente agronómico", "ui_chat.png", "Asistente conversacional Agroplaga AI con contexto de parcela y diagnóstico por imagen.", "• Consultas sobre riego, clima, plagas, poda y manejo integrado.\n• Adjuntar una fotografía para obtener un diagnóstico.\n• Rechazo de consultas fuera de la agronomía del melocotón."),
]
for title, img, cap, inter in UI:
    ui = par_after(ui, title, style=H3)
    ui = par_after(ui, "Interacciones esperadas:", bold=True)
    ui = par_after(ui, inter)
    ui = picture_after(ui, os.path.join(FIG, img), max_w=5.6, max_h=6.6)
    ui = caption_after(ui, "Figura", cap + " (Fuente: elaboración propia)")

# ------------------------------------------------------------------ 8. capitulo 6
set_text(p(486), "Capítulo 6 – Desarrollo de la solución", style=H1, align=CENTER)
set_text(p(487), "6.1 Implementación", style=H2)
set_text(p(488), "Para dar solución al problema planteado se construyó Agroplaga AI, una aplicación web orientada al monitoreo ambiental, al "
         "video en vivo de la parcela y al apoyo en el manejo de plagas en plantaciones de melocotón. El sistema permite consultar las lecturas "
         "recientes del sensor DHT22, observar el cultivo a través del ESP32-CAM, analizar imágenes con inteligencia artificial, registrar "
         "parcelas, administrar fichas de plagas, documentar actividades agrícolas y utilizar un asistente conversacional especializado en "
         "agronomía del melocotón.")
set_text(p(489), "El prototipo se desarrolló y validó inicialmente en ambiente local, utilizando un frontend en React con Vite, un backend en "
         "Node.js con Express y una base de datos MySQL, junto con un simulador del ESP32-CAM y luego el nodo físico. Esta arquitectura "
         "permitió comprobar la comunicación entre el nodo de campo, la interfaz, la API, la base de datos y el servicio de inteligencia "
         "artificial antes de considerar un despliegue en producción.")
set_text(p(490), "La aplicación incluye módulos para el panel de monitoreo ambiental, cámara en vivo con análisis por IA, gestión de plagas, "
         "registro de nuevas plagas, administración de parcelas, seguimiento de actividades agrícolas, exportación de historial a CSV y "
         "chatbot agronómico. El panel principal presenta temperatura y humedad por parcela junto con la última detección de la cámara, "
         "mientras que el backend normaliza las lecturas, evalúa umbrales, retransmite el video y expone endpoints para consultar la "
         "información más reciente.")
set_text(p(493), "6.1.1 Tecnologías utilizadas", style=H3)
set_text(p(494), "Frontend", bold=True)
tech = chain(p(494), [
    ("• JavaScript, React 19 y Vite; React Router para la navegación; CSS por módulo; Recharts para gráficos.", N),
    ("Backend y almacenamiento", N, {"bold": True}),
    ("• Node.js con Express 5 como API REST; Sequelize sobre MySQL 8; Multer para la carga de imágenes; retransmisión de video MJPEG mediante multipart/x-mixed-replace.", N),
    ("Inteligencia artificial", N, {"bold": True}),
    ("• Servicio de visión de OpenAI para el diagnóstico de imágenes con salida JSON estructurada (plaga, confianza, severidad, recomendaciones); Ollama como alternativa local para el chat de texto.", N),
    ("Hardware del nodo de campo", N, {"bold": True}),
    ("• ESP32-CAM AI-Thinker (OV2640, Wi-Fi), sensor DHT22, panel solar de 6 V, módulo de carga TP4056, batería 18650 y elevador de 5 V; firmware en Arduino con las librerías esp_camera, esp_http_server y DHT sensor library.", N),
    ("Herramientas", N, {"bold": True}),
    ("• Visual Studio Code, Arduino IDE, Docker (MySQL), Postman, jest y supertest para pruebas, Git y GitHub.", N),
])
hw = par_after(tech, "6.1.2 Nodo de campo: ESP32-CAM, DHT22 y punto de carga solar", style=H3)
hw = chain(hw, [
    ("El nodo de campo concentra en un solo módulo ESP32-CAM la captura de video, la lectura del sensor DHT22 y la comunicación Wi-Fi. El "
     "sensor se conecta al pin GPIO 13 con una resistencia pull-up de 10 kΩ, y el módulo se alimenta a 5 V desde un punto de carga solar "
     "formado por un panel de 6 V, un módulo TP4056 con protección, una batería Li-ion 18650 y un elevador de voltaje. La Figura siguiente "
     "muestra el esquema de conexión.", N),
])
hw = picture_after(hw, os.path.join(FIG, "hardware.png"))
hw = caption_after(hw, "Figura", "Esquema de conexión del nodo de campo con punto de carga solar. (Fuente: elaboración propia)")
hw = par_after(hw, "El firmware, desarrollado en Arduino, expone los mismos servicios que el CameraWebServer de Espressif: una fotografía JPEG "
               "en /capture y un flujo MJPEG continuo en el puerto 81 (/stream). Cada 60 segundos lee el DHT22 y envía la temperatura y la "
               "humedad al backend mediante POST /api/readings, donde se evalúan los umbrales de la parcela. Opcionalmente puede enviar "
               "fotografías al servidor (modo push) cuando este no puede alcanzar la dirección IP de la cámara.")
hw = picture_after(hw, os.path.join(FIG, "cod_firmware_dht.png"))
hw = caption_after(hw, "Figura", "Fragmento del firmware del ESP32-CAM: lectura del DHT22 y envío de lecturas al backend. (Fuente: elaboración propia)")
hw = picture_after(hw, os.path.join(FIG, "cod_firmware_stream.png"))
hw = caption_after(hw, "Figura", "Fragmento del firmware del ESP32-CAM: servidor de video MJPEG en el puerto 81. (Fuente: elaboración propia)")
set_text(p(497), "6.2 Desarrollo de la aplicación", style=H2)
replace_picture(p(516), os.path.join(FIG, "ui_dashboard.png"), max_w=5.6, max_h=6.2)
set_text(p(518), "Figura 20. Panel de monitoreo ambiental con tarjetas de temperatura y humedad del DHT22, última detección de cámara y alertas inteligentes (elaboración propia).")
cam = par_after(p(559), "6.2.1 Módulo de cámara en vivo y detección automática de plagas", style=H3)
cam = chain(cam, [
    ("El módulo de cámara resuelve dos limitaciones del ESP32-CAM: solo admite uno o dos espectadores y no puede analizar imágenes por sí "
     "mismo. El backend mantiene un concentrador de cuadros (FrameHub) por cámara: abre una única conexión al flujo MJPEG del nodo cuando "
     "hay al menos un espectador, trocea el flujo localizando los marcadores de inicio y fin de cada JPEG, conserva en memoria el último "
     "cuadro y lo retransmite a todos los navegadores conectados mediante multipart/x-mixed-replace. Cuando el último espectador cierra la "
     "página, la conexión con el nodo se libera automáticamente.", N),
])
cam = picture_after(cam, os.path.join(FIG, "cod_mjpeg.png"))
cam = caption_after(cam, "Figura", "Parser incremental del flujo MJPEG que separa los cuadros JPEG por sus marcadores SOI y EOI (elaboración propia).")
cam = picture_after(cam, os.path.join(FIG, "cod_stream_route.png"))
cam = caption_after(cam, "Figura", "Endpoint GET /api/camaras/:id/stream que retransmite el video a los navegadores descartando cuadros para clientes lentos (elaboración propia).")
cam = par_after(cam, "Desde la página Cámara el agricultor observa el video, captura una fotografía o pulsa «Analizar con IA». El backend toma el "
                "último cuadro, lo guarda en disco, consulta el contexto de la parcela (umbrales, últimas lecturas y alerta más reciente) y "
                "envía la imagen junto con un prompt estructurado al servicio de visión, que responde un JSON con la plaga detectada, la "
                "confianza, la severidad, un resumen y recomendaciones de manejo integrado. La respuesta se normaliza y se registra en la "
                "tabla detecciones, de modo que aparece en el historial de la cámara, en el panel principal y, si requiere acción, en las "
                "alertas inteligentes.")
cam = picture_after(cam, os.path.join(FIG, "ui_camara.png"), max_w=5.2, max_h=7.2)
cam = caption_after(cam, "Figura", "Página de cámara en vivo mostrando el video del nodo, el estado en línea y los controles de captura y análisis (elaboración propia).")
cam = picture_after(cam, os.path.join(FIG, "cod_vision_prompt.png"))
cam = caption_after(cam, "Figura", "Construcción del prompt estructurado con el contexto de la parcela para el modelo de visión (elaboración propia).")
cam = par_after(cam, "El análisis automático se configura por cámara con un intervalo mínimo de 60 segundos. Un planificador del servidor "
                "revisa periódicamente las cámaras activas, omite el ciclo si no hay un cuadro reciente y aplica una espera creciente cuando "
                "el servicio de IA falla, con lo que se controla el costo del servicio externo.")
cam = picture_after(cam, os.path.join(FIG, "cod_scheduler.png"))
cam = caption_after(cam, "Figura", "Planificador de análisis automático con intervalo mínimo, verificación de cuadro fresco y espera ante fallos (elaboración propia).")
set_text(p(561), "6.3 Pruebas de Agroplaga AI", style=H2)
set_text(p(578), "Flujo de sensores: el nodo ESP32-CAM lee el DHT22 y envía la lectura al backend, el backend valida la información, la almacena y evalúa los umbrales de la parcela.")
tst = par_after(p(614), "6.3.4 Pruebas automatizadas", style=H3)
tst = par_after(tst, "Además de las pruebas manuales, el backend cuenta con una suite automatizada ejecutada con jest y supertest que cubre "
                "el parser MJPEG, la normalización del diagnóstico de IA, el registro de lecturas, el chat con el servicio de IA simulado y el "
                "flujo completo de cámara en modo push: creación de la cámara, envío de un cuadro, obtención de la fotografía, análisis con la "
                "IA simulada mediante nock, consulta de la detección y eliminación en cascada. Las pruebas de sistema verifican que la "
                "detección queda almacenada en MySQL. La suite se compone de 51 pruebas en 7 archivos y se ejecuta con el comando npm test.")
tst = picture_after(tst, os.path.join(FIG, "pruebas_jest.png"), max_h=7.2)
tst = caption_after(tst, "Figura", "Resultado de la suite de pruebas automatizadas del backend (elaboración propia).")
docs = par_after(tst, "6.4 Documentación", style=H2)
docs = chain(docs, [
    ("6.4.1 Manual de usuario", H3),
    ("El manual de usuario describe cómo registrar parcelas y cámaras, interpretar el panel ambiental, observar el video en vivo, capturar "
     "y analizar imágenes, leer el resultado del diagnóstico según su severidad, activar el análisis automático, registrar actividades y "
     "consultar al asistente agronómico. Se encuentra en el repositorio del proyecto en el archivo PGAPAYBABACK/CAMARA_ESP32.md junto con la "
     "guía de uso del simulador.", N),
    ("6.4.2 Manual técnico", H3),
    ("El manual técnico está dirigido a quien instale o mantenga el sistema. Incluye la instalación del backend y la base de datos "
     "(PGAPAYBABACK/LEVANTAR_BACKEND.md), la migración de la base de datos para cámaras y detecciones, las variables de entorno, la "
     "descripción de los endpoints de la API, el armado del nodo de campo y la carga del firmware en el ESP32-CAM (firmware/esp32cam/README.md), "
     "y la ejecución de las pruebas automatizadas.", N),
    ("6.5 Aplicación", H2),
    ("El código fuente completo de Agroplaga AI (frontend, backend, firmware del ESP32-CAM, simulador y pruebas) se encuentra en el "
     f"repositorio público del proyecto: {REPO}. Para ejecutarlo localmente se levanta MySQL, se instala el backend y el frontend con npm "
     "install y se inician con npm start y npm run dev; el simulador de cámara se ejecuta con npm run simular-camara.", N),
])

# encabezados vacios que ensucian el indice
for par in P:
    if par.style.name.startswith("Heading") and not par.text.strip():
        par.style = doc.styles[N]

# ------------------------------------------------------------------ 9. conclusiones / recomendaciones
def wipe_empty_after(start_idx, stop_idx):
    for i in range(start_idx + 1, stop_idx):
        par = p(i)
        if not par.text.strip() and not par._p.findall(".//" + qn("w:drawing")):
            delete_par(par)
wipe_empty_after(635, 675); wipe_empty_after(675, 715)
CONCL = [
    "a) Se desarrolló Agroplaga AI, un sistema que integra un nodo de campo de bajo costo, una aplicación web y un servicio de inteligencia "
    "artificial para apoyar la detección temprana de plagas en cultivos de melocotón de FRUTAGRU. El sistema permite observar el cultivo en "
    "tiempo real, medir sus condiciones ambientales, diagnosticar imágenes y registrar el manejo agrícola desde cualquier navegador.",
    "b) La combinación de un módulo ESP32-CAM con un sensor DHT22 demostró ser suficiente para cubrir el monitoreo visual y ambiental de una "
    "parcela con un solo dispositivo, y el punto de carga solar permite instalarlo sin acceso a la red eléctrica. Concentrar el video en el "
    "servidor resolvió la limitación del ESP32-CAM de atender solo uno o dos espectadores.",
    "c) El uso de un modelo multimodal de visión accesible mediante API permitió obtener diagnósticos estructurados con plaga, confianza, "
    "severidad y recomendaciones sin entrenar un modelo propio ni recolectar un conjunto de imágenes etiquetadas, lo que reduce el tiempo y "
    "el costo de implementación. Enviar al modelo el contexto ambiental de la parcela mejora la pertinencia de las recomendaciones.",
    "d) Las alertas basadas en umbrales de temperatura y humedad, sumadas a las detecciones de la cámara que requieren acción, entregan al "
    "agricultor información oportuna y comprensible para decidir cuándo intervenir y evitar aplicaciones innecesarias de pesticidas.",
    "e) Las pruebas unitarias, de integración, de sistema y de validación, junto con la suite automatizada de 51 pruebas, comprobaron la "
    "comunicación entre el nodo, el backend, la base de datos, el servicio de IA y la interfaz, así como el comportamiento del sistema ante "
    "entradas inválidas y ante la ausencia de la cámara.",
    "f) Agroplaga AI constituye una base funcional y documentada que puede ampliarse a más parcelas y cultivos. La precisión del diagnóstico "
    "depende de la calidad de la imagen y de las condiciones de iluminación, por lo que el sistema debe entenderse como un apoyo a la "
    "decisión del agricultor y del técnico, no como un sustituto de la inspección en campo.",
]
cur = p(635)
for t in CONCL: cur = par_after(cur, t)
RECOM = [
    "a) Se recomienda instalar el nodo de prueba en una parcela de FRUTAGRU durante al menos un ciclo productivo, ajustando la ubicación y "
    "el enfoque de la cámara hacia el follaje y los frutos, y registrando la retroalimentación de los agricultores para mejorar la interfaz.",
    "b) Conviene calibrar los umbrales de temperatura y humedad de cada parcela con el apoyo de un técnico agrónomo y comparar las "
    "detecciones de la IA con inspecciones en campo, para medir la precisión real del diagnóstico y ajustar el intervalo del análisis automático.",
    "c) Se recomienda dimensionar el punto de carga solar según la zona: un panel de mayor potencia o una segunda batería permitirán "
    "mantener el video disponible durante más horas en temporadas nubladas, y una caja estanca protegerá el nodo de la lluvia.",
    "d) Para ampliar la cobertura se sugiere añadir más nodos ESP32-CAM registrándolos desde la aplicación, y evaluar el modo push cuando el "
    "servidor no pueda alcanzar la red de la parcela. También se recomienda incorporar autenticación de usuarios antes de un despliegue abierto.",
    "e) Es conveniente desplegar el backend y la base de datos en un servidor accesible desde internet, configurar HTTPS y respaldos, y "
    "monitorear el consumo del servicio de IA para mantener el costo operativo bajo control.",
    "f) Por último, se recomienda mantener actualizadas las dependencias de Node.js, React y el firmware del ESP32, ejecutar la suite de "
    "pruebas automatizadas antes de cada cambio y conservar la documentación técnica del repositorio al día.",
]
cur = p(675)
for t in RECOM: cur = par_after(cur, t)

# ------------------------------------------------------------------ 10. glosario
for i in range(716, 751): delete_par(p(i))
GLOS = [
    ("A", None), ("Alerta por umbral", "Aviso que genera el sistema cuando una lectura ambiental queda fuera del rango mínimo o máximo configurado para la parcela."),
    ("API (Interfaz de Programación de Aplicaciones)", "Conjunto de endpoints HTTP que permiten a la aplicación web, al nodo de campo y a otros programas intercambiar datos con el servidor."),
    ("B", None), ("Base de datos", "Conjunto estructurado de información, en este caso MySQL, donde se almacenan parcelas, lecturas, alertas, plagas, cámaras y detecciones."),
    ("D", None), ("Detección", "Registro generado por el análisis de una imagen con inteligencia artificial, que incluye la plaga identificada, la confianza, la severidad, un resumen y recomendaciones."),
    ("DHT22", "Sensor digital de temperatura y humedad del aire utilizado por el nodo de campo."),
    ("E", None), ("ESP32-CAM", "Módulo de bajo costo que integra el microcontrolador ESP32 con Wi-Fi y una cámara OV2640; constituye el nodo de campo de Agroplaga AI."),
    ("F", None), ("Firmware", "Programa grabado en el ESP32-CAM que controla la cámara, lee el sensor y se comunica con el servidor."),
    ("FrameHub", "Componente del servidor que conserva el último cuadro de cada cámara y lo distribuye a los navegadores y al módulo de análisis."),
    ("I", None), ("Inteligencia Artificial (IA)", "Área de la informática que desarrolla sistemas capaces de realizar tareas que normalmente requieren inteligencia humana, como interpretar imágenes."),
    ("Internet de las Cosas (IoT)", "Interconexión de dispositivos con sensores que recopilan y transmiten datos del entorno, como el nodo ESP32-CAM con DHT22."),
    ("M", None), ("Manejo Integrado de Plagas (MIP)", "Estrategia que combina prácticas culturales, biológicas, físicas y químicas racionales para minimizar el impacto de las plagas."),
    ("MJPEG", "Formato de video en el que cada cuadro es una imagen JPEG independiente; es el que sirve el ESP32-CAM y retransmite el servidor."),
    ("Modelo multimodal de visión", "Modelo de inteligencia artificial capaz de interpretar imágenes y texto simultáneamente y de responder en lenguaje natural."),
    ("P", None), ("Punto de carga solar", "Conjunto de panel fotovoltaico, módulo de carga, batería y regulador que alimenta el nodo de campo de forma autónoma."),
    ("R", None), ("REST", "Estilo de diseño de APIs basado en recursos y métodos HTTP (GET, POST, PUT, DELETE) empleado por el backend."),
    ("S", None), ("Severidad", "Nivel de gravedad asignado a una detección (ninguna, baja, media o alta) que determina si requiere acción."),
    ("Simulador de cámara", "Programa que imita el comportamiento del ESP32-CAM sirviendo imágenes de prueba, usado para desarrollar y probar sin hardware."),
    ("U", None), ("Umbral", "Valor mínimo y máximo aceptable de una variable ambiental, configurado por parcela y sensor."),
    ("UX (Experiencia de Usuario)", "Percepción y facilidad de uso que tiene el agricultor al interactuar con la aplicación."),
]
cur = p(715)
for term, desc in GLOS:
    cur = par_after(cur, term if desc is None else f"{term}: {desc}", bold=(desc is None))

# ------------------------------------------------------------------ 11. referencias APA 7
for i in range(753, len(P) + 1): delete_par(p(i))
REFS = [
    "ANAPDE. (2019). Informe técnico anual. Asociación Nacional del Programa de Desarrollo de Exportación.",
    "Aosong Electronics. (2015). Digital-output relative humidity & temperature sensor/module DHT22 (AM2302) datasheet. Aosong Electronics Co., Ltd.",
    "Arias, F. G. (2012). El proyecto de investigación: Introducción a la metodología científica (6.ª ed.). Episteme.",
    "Arias, L. (2012). Innovación agrícola en Centroamérica. Editorial Agrotech.",
    "Cruz, P., & López, M. (2017). Trampas feromónicas para Grapholita molesta. Revista de Fitopatología, 22(3), 45–58.",
    "Delgado, R., & Morales, S. (2013). Trampas cromáticas para moscas de la fruta. Acta Agronómica, 60(2), 101–110.",
    "Espressif Systems. (2023). ESP32 series datasheet. Espressif Systems. https://www.espressif.com/en/support/documents/technical-documents",
    "FAO. (2020). Agricultura en cifras. Organización de las Naciones Unidas para la Alimentación y la Agricultura.",
    "Fernández, J. (2020). Aplicación de los sensores de IoT en la gestión de los pesticidas agrícolas. Revista Agrotecnológica, 12(2), 55–70.",
    "FRUTAGRU. (2021). Memoria institucional 2021. Asociación de Fruticultores de Guatemala.",
    "García, J., & Rivera, L. (2015). Umbrales económicos en frutales. Estudios Agrícolas, 18(4), 77–89.",
    "Goodfellow, I., Bengio, Y., & Courville, A. (2016). Deep learning. MIT Press.",
    "Hernández, P., & Torres, G. (2016). Costos de producción de frutales en Guatemala. Revista Económica Agrícola, 8(2), 33–47.",
    "Hernández Sampieri, R., Fernández Collado, C., & Baptista Lucio, P. (2014). Metodología de la investigación (6.ª ed.). McGraw-Hill.",
    "ICTA. (2020). Estudios sobre el uso de pesticidas en Guatemala. Instituto de Ciencia y Tecnología Agrícolas.",
    "INTECA. (2019). API de pronóstico del tiempo para Guatemala. Instituto Nacional de Tecnología Agropecuaria.",
    "Jiménez, R., Castillo, M., & Ortiz, D. (2018). Biología de Anastrepha spp. en Centroamérica. Entomología Tropical, 40(1), 23–39.",
    "Kamilaris, A., & Prenafeta-Boldú, F. X. (2018). Deep learning in agriculture: A survey. Computers and Electronics in Agriculture, 147, 70–90. https://doi.org/10.1016/j.compag.2018.02.016",
    "Kumar, S., & Lee, J. (2019). Evaluación de sensores DHT22 vs. SHT31. Sensors Journal, 19(3), 567–575.",
    "Lehrig, S. (2021). Progressive web apps en comunidades rurales. Digital Agriculture Review, 3(1), 34–48.",
    "Li, X., Wang, H., & Chen, Y. (2020). Sistemas embebidos con TensorFlow Lite. Embedded Systems Journal, 8(4), 120–130.",
    "López, A. (2022). Uso de plataformas digitales en gestión agrícola. Journal of Precision Agriculture, 9(1), 12–25.",
    "López, A., & Gómez, B. (2018). Tendencias de diversificación agrícola en Centroamérica. Revista Centroamericana de Desarrollo Rural, 6(2), 45–60.",
    "MAGA. (2020). Guía de MIP para frutales. Ministerio de Agricultura, Ganadería y Alimentación.",
    "MAGA. (2021). Informe anual de productividad agrícola. Ministerio de Agricultura, Ganadería y Alimentación.",
    "Martínez, F., & Salazar, D. (2020). Mercado y cadenas de valor del melocotón. Economía y Desarrollo Rural, 15(3), 110–125.",
    "Meta Platforms. (2024). React: The library for web and native user interfaces. https://react.dev/",
    "OIRSA. (2017). Manual de muestreo de plagas cuarentenarias. Organismo Internacional Regional de Sanidad Agropecuaria.",
    "OpenAI. (2024). OpenAI API documentation: Vision. https://platform.openai.com/docs/guides/vision",
    "OpenJS Foundation. (2024). Node.js documentation. https://nodejs.org/docs/",
    "Oracle Corporation. (2024). MySQL 8.0 reference manual. https://dev.mysql.com/doc/refman/8.0/en/",
    "Organización de las Naciones Unidas para la Alimentación y la Agricultura. (2021). Innovaciones digitales para la agricultura familiar. FAO.",
    "Pérez, L., Ramírez, A., & Méndez, J. (2019). Rendimiento y rentabilidad del melocotón. Revista Agrícola de Guatemala, 25(2), 88–102.",
    "Ramírez, E., & Chávez, V. (2016). Resistencia de plagas a acaricidas en Guatemala. Revista Fitotécnica, 28(1), 15–27.",
    "Ramírez, H., & Gómez, P. (2021). Lecciones aprendidas de AgroIA en Colombia. Revista Colombiana de Tecnología Agrícola, 14(2), 60–74.",
    "Ramírez, J., Soto, L., & Paz, M. (2021). Sistemas de alerta temprana en web. Journal of Agricultural Informatics, 11(1), 67–82.",
    "Rodríguez, P. P. (2021). El impacto de Internet de las cosas en la prevención de las plagas. Agrotech Magazine, 28(4), 70–78.",
    "Ruiz, M., Díaz, C., & Rojas, F. (2021). IoT como servicio para pequeños productores. Agritech Business, 6(2), 45–60.",
    "Russell, S., & Norvig, P. (2016). Artificial intelligence: A modern approach (3.ª ed.). Pearson.",
    "Silva, M., & Ferreira, J. (2019). Análisis costo-beneficio de IoT en frutales. International Journal of Agricultural Technology, 15(4), 210–224.",
    "Smith, D., & Jones, E. (2020). Cinco años de MYIPM: Lecciones aprendidas. Journal of Extension, 58(6).",
    "Soto, R., Mejía, A., & Cano, J. (2014). Contaminación de fuentes hídricas por pesticidas. Hidrogeoquímica, 9(2), 101–115.",
    "Stern, V. M., Smith, R. F., van den Bosch, R., & Hagen, K. S. (1959). The integration of chemical and biological control of the spotted alfalfa aphid. Hilgardia, 29(2), 81–101.",
    "Vargas, E., Molina, R., & Pineda, S. (2015). Control biológico de pulgones en frutales. Revista Colombiana de Entomología, 41(2), 89–98.",
    "Wolfert, S., Ge, L., Verdouw, C., & Bogaardt, M.-J. (2017). Big data in smart farming: A review. Agricultural Systems, 153, 69–80. https://doi.org/10.1016/j.agsy.2017.01.023",
    "Zhang, Q., Liu, Y., & Gong, C. (2019). CNN para detección de plagas en manzano. Computers and Electronics in Agriculture, 162, 434–443.",
]
def sort_key(r): return re.sub(r"[^a-záéíóúñ ]", "", r.lower())
cur = p(752)
for r in sorted(REFS, key=sort_key):
    cur = par_after(cur, r)
    cur.paragraph_format.left_indent = Inches(0.5); cur.paragraph_format.first_line_indent = Inches(-0.5)

# ------------------------------------------------------------------ 12. pasada final: leyendas con campo SEQ
CAP_RE = re.compile(r"^(Figura|Tabla)\s*(\d+)\s*[.:]?\s*(.*)$", re.S)
counts = {"Figura": 0, "Tabla": 0}
def build_caption(par, kind, text):
    counts[kind] += 1
    had_pb = clear_runs(par); par.alignment = CENTER
    text = CTRL_RE.sub("", text)
    r = par.add_run(f"{kind} "); r.italic = True; r.font.name = "Times New Roman"; r.font.size = Pt(12)
    fs = OxmlElement("w:fldSimple"); fs.set(qn("w:instr"), f" SEQ {kind} \\* ARABIC ")
    rr = OxmlElement("w:r"); rpr = OxmlElement("w:rPr"); i_ = OxmlElement("w:i"); rpr.append(i_); rr.append(rpr)
    t = OxmlElement("w:t"); t.text = str(counts[kind]); rr.append(t); fs.append(rr); par._p.append(fs)
    text = text.strip()
    text = re.sub(r"\(elaboraci[oó]n propia\)\.?$", "(Fuente: elaboración propia).", text, flags=re.I)
    text = re.sub(r"\(Fuente:\s*elaboraci[oó]n propia\)\.?$", "(Fuente: elaboración propia).", text, flags=re.I)
    if not text.endswith("."): text += "."
    r2 = par.add_run(f". {text}"); r2.italic = True; r2.font.name = "Times New Roman"; r2.font.size = Pt(12)
    par.paragraph_format.keep_with_next = False
    if had_pb: page_break_after(par)

def ptext(par):
    """Texto leido directamente de los nodos w:t (los parrafos creados en memoria no exponen .text hasta recargar)."""
    return "".join(t.text or "" for t in par._p.iter(qn("w:t")))
all_pars = [Paragraph(el, doc._body) for el in doc.element.body.iterchildren() if el.tag == qn("w:p")]
print("parrafos en pasada final:", len(all_pars), "| marcados:", sum(1 for q in all_pars if ptext(q).startswith(SEQ_MARK)))
for par in all_pars:
    txt = ptext(par)
    if txt.startswith(SEQ_MARK):
        _, kind, text = txt.split("|", 2)
        build_caption(par, kind, text); continue
    m = CAP_RE.match(txt.strip())
    if m and len(txt) < 400 and not par.style.name.startswith("Heading"):
        build_caption(par, m.group(1), m.group(3))
print("Leyendas:", counts)

# ------------------------------------------------------------------ 13. estilos APA, numeros de pagina, actualizar campos
doc.styles[H3].font.italic = True
for st in (H1, H2, H3, H4): doc.styles[st].font.name = "Times New Roman"; doc.styles[st].font.size = Pt(12); doc.styles[st].font.bold = True
for si, sec in enumerate(doc.sections):
    sec.header.is_linked_to_previous = False
    hp = sec.header.paragraphs[0] if sec.header.paragraphs else sec.header.add_paragraph()
    clear_runs(hp); hp.alignment = RIGHT; add_field(hp, "PAGE", "1")
    if si == 0:
        sec.different_first_page_header_footer = True
        for fp in sec.first_page_header.paragraphs: clear_runs(fp)
uf = OxmlElement("w:updateFields"); uf.set(qn("w:val"), "true"); doc.settings.element.append(uf)
rest = [(i + 1, q.text[:80]) for i, q in enumerate(doc.paragraphs) if re.search(r"\w\?\w|(^|\s)\?\w", q.text) and "http" not in q.text]
print("Mojibake restante:", rest)
doc.save(DST)
print("guardado:", DST)
